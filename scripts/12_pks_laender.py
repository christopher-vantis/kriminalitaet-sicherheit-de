#!/usr/bin/env python3
"""
12_pks_laender.py — Kriminalität je Bundesland (PKS-Länderdaten, Berichtsjahr 2025)
===================================================================================
Quelle: data/raw/pks/laender/2025/LA-F-02-T01-Laender-Faelle-HZ_xls.xlsx
(Tabelle 01 „Fälle", Bereich Länder; Bundesland, erfasste Fälle, Häufigkeitszahl
auf Basis Zensus 2022, Versuche, aufgeklärte Fälle, Aufklärungsquote).

Ausgabe: output/pks_laender_2025.csv
  bundesland; schluessel; delikt; faelle; hz; versuche; aufgeklaert; aq
"""
import csv
import pathlib

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
QUELLE = ROOT / "data/raw/pks/laender/2025/LA-F-02-T01-Laender-Faelle-HZ_xls.xlsx"
OUT = ROOT / "output/pks_laender_2025.csv"

# Deliktsgruppen, die in die App gehen
GEWUENSCHT = {
    "------": "Straftaten insgesamt",
    "****00": "Diebstahl (gesamt)",
    "435*00": "Wohnungseinbruchdiebstahl",
    "210000": "Raub",
    "222000": "Gefährliche Körperverletzung",
    "100000": "Straftaten gegen die sexuelle Selbstbestimmung",
    "200000": "Rohheitsdelikte und Straftaten gegen die persönliche Freiheit",
    "500000": "Vermögens- und Fälschungsdelikte",
    "674000": "Sachbeschädigung",
    "700000": "Straftaten gegen strafrechtliche Nebengesetze",
}

LAENDER = {"Baden-Württemberg", "Bayern", "Berlin", "Brandenburg", "Bremen", "Hamburg",
           "Hessen", "Mecklenburg-Vorpommern", "Niedersachsen", "Nordrhein-Westfalen",
           "Rheinland-Pfalz", "Saarland", "Sachsen", "Sachsen-Anhalt",
           "Schleswig-Holstein", "Thüringen"}


def zahl(x):
    if x is None:
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def main():
    wb = openpyxl.load_workbook(QUELLE, read_only=True)
    ws = wb[wb.sheetnames[0]]
    zeilen = []
    schluessel = delikt = None
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=10, values_only=True):
        r = list(row) + [None] * 10
        s, d, bl = r[0], r[1], r[2]
        if s and str(s).strip() and str(s).strip() != "1":
            schluessel = str(s).strip()
        if d and str(d).strip() and str(d).strip() != "2":
            delikt = str(d).strip()
        if not bl or str(bl).strip() not in LAENDER:
            continue
        if schluessel not in GEWUENSCHT:
            continue
        zeilen.append(dict(
            bundesland=str(bl).strip(), schluessel=schluessel,
            delikt=GEWUENSCHT[schluessel],
            faelle=zahl(r[3]), hz=zahl(r[4]), versuche=zahl(r[5]),
            aufgeklaert=zahl(r[7]), aq=zahl(r[8]),
        ))
    OUT.parent.mkdir(exist_ok=True)
    with OUT.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["bundesland", "schluessel", "delikt",
                                           "faelle", "hz", "versuche",
                                           "aufgeklaert", "aq"], delimiter=";")
        w.writeheader()
        w.writerows(zeilen)

    print(f"{len(zeilen)} Zeilen -> {OUT}")
    print("\nStraftaten insgesamt je Bundesland (Fälle, HZ, AQ):")
    ges = sorted([z for z in zeilen if z["schluessel"] == "------"],
                 key=lambda z: -(z["hz"] or 0))
    for z in ges:
        print(f"  {z['bundesland']:26s} {z['faelle']:>9,.0f} Fälle   HZ {z['hz']:>7.1f}   "
              f"AQ {z['aq']:>5.1f} %")
    print(f"\nDelikte abgedeckt: {len(set(z['schluessel'] for z in zeilen))}")
    print(f"Länder: {len(set(z['bundesland'] for z in zeilen))}")


if __name__ == "__main__":
    main()
